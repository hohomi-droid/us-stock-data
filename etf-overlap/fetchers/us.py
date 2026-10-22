"""미국 운용사별 일일 보유종목 수집기.

각 함수는 dict 반환:
  {"asof": "YYYY-MM-DD", "units": float|None,
   "holdings": {SYMBOL: {"name", "shares", "weight", "type", "yahoo"}}}
SYMBOL 은 비교 키(미국 티커 또는 거래소 접미사 포함), yahoo 는 가격 조회용 심볼.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime

import requests

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36",
      "Accept": "*/*"}
S = requests.Session()
S.headers.update(UA)


def num(x):
    if x is None:
        return None
    s = str(x).replace(",", "").replace("$", "").replace("%", "").strip()
    if s in ("", "-", "--", "N/A"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def iso(s: str) -> str | None:
    s = (s or "").strip().strip('"')
    for fmt in ("%m/%d/%Y", "%b %d, %Y", "%Y-%m-%d", "%m-%d-%Y", "%d-%b-%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None


# 해외 상장 종목: 운용사 표기 → Yahoo 접미사
EXCH_SUFFIX = {
    "Tokyo Stock Exchange": ".T", "London Stock Exchange": ".L", "Hong Kong Exchanges And Clearing Ltd": ".HK",
    "Euronext Amsterdam": ".AS", "Nyse Euronext - Euronext Paris": ".PA", "Xetra": ".DE", "Deutsche Boerse Xetra": ".DE",
    "SIX Swiss Exchange": ".SW", "Korea Exchange": ".KS", "Taiwan Stock Exchange": ".TW", "Toronto Stock Exchange": ".TO",
    "Borsa Italiana": ".MI", "Bolsa De Madrid": ".MC", "Nasdaq Omx Nordic": ".ST", "Asx - All Markets": ".AX",
    "Nasdaq Copenhagen": ".CO", "Oslo Bors Asa": ".OL", "National Stock Exchange Of India": ".NS",
}
US_EXCH = ("NASDAQ", "New York Stock Exchange Inc.", "Nyse Mkt Llc", "Cboe BZX", "NYSE Arca", "Nyse Arca")


def yahoo_sym(ticker: str, exchange: str | None = None, country: str | None = None) -> str:
    t = (ticker or "").strip().upper().replace("/", "-").replace(" ", "-")
    if exchange and exchange not in US_EXCH:
        for k, suf in EXCH_SUFFIX.items():
            if k.lower() in exchange.lower():
                if suf in (".HK",):
                    t = t.lstrip("0").zfill(4)
                return t + suf
    if t.endswith(".") or t.endswith("-"):
        t = t[:-1]
    return t.replace(".", "-")  # BRK.B → BRK-B


def is_equity(asset_class: str | None) -> bool:
    a = (asset_class or "").lower()
    return a in ("equity", "stock", "common stock", "equities", "ordinary shares", "adr", "") or "equit" in a or "common" in a


# ---------------------------------------------------------------- iShares
def ishares(product_id: str, slug: str):
    url = f"https://www.ishares.com/us/products/{product_id}/{slug}/latest-holdings.csv"
    r = S.get(url, timeout=30)
    r.raise_for_status()
    text = r.content.decode("utf-8-sig", errors="replace")
    lines = text.splitlines()
    asof, units, start = None, None, None
    for i, ln in enumerate(lines[:20]):
        if ln.startswith("Fund Holdings as of"):
            asof = iso(ln.split(",", 1)[1])
        elif ln.startswith("Shares Outstanding"):
            units = num(ln.split(",", 1)[1].strip().strip('"'))
        elif ln.startswith("Ticker,"):
            start = i
            break
    if start is None:
        raise ValueError("iShares CSV 헤더를 찾지 못함")
    body = "\n".join(lines[start:])
    out = {}
    for row in csv.DictReader(io.StringIO(body)):
        tk = (row.get("Ticker") or "").strip()
        if not tk or tk.startswith("\xa0") or row.get("Name") is None:
            continue
        if not is_equity(row.get("Asset Class")):
            continue
        sh = num(row.get("Quantity") or row.get("Shares"))
        if not sh:
            continue
        y = yahoo_sym(tk, row.get("Exchange"), row.get("Location"))
        out[y] = {"name": (row.get("Name") or "").title(), "shares": sh,
                  "weight": num(row.get("Weight (%)")) or 0.0, "type": "equity", "yahoo": y}
    return {"asof": asof, "units": units, "holdings": out}


def ishares_weights(product_id: str, slug: str):
    """벤치마크 대용 지수 ETF 비중 {yahoo_sym: weight%}"""
    d = ishares(product_id, slug)
    return {s: h["weight"] for s, h in d["holdings"].items()}, d["asof"]


# ---------------------------------------------------------------- JPMorgan
JPM_BASE = "https://am.jpmorgan.com/FundsMarketingHandler"


def jpm(cusip: str):
    import openpyxl
    q = f"type=dailyETFHoldings&cusip={cusip}&country=us&role=adv&locale=en-US"
    r = S.get(f"{JPM_BASE}/excel?{q}", timeout=40)
    r.raise_for_status()
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, rw in enumerate(rows) if rw and any(str(c).strip().lower() == "ticker" for c in rw if c))
    hdr = [str(c).strip() if c else "" for c in rows[hdr_i]]
    col = {h.lower(): i for i, h in enumerate(hdr)}

    def g(rw, *names):
        for n in names:
            if n in col and col[n] < len(rw):
                return rw[col[n]]
        return None
    asof = None
    for rw in rows[:hdr_i]:
        for c in rw or ():
            if c and re.search(r"\d{1,2}/\d{1,2}/\d{4}", str(c)):
                asof = iso(re.search(r"\d{1,2}/\d{1,2}/\d{4}", str(c)).group(0))
    out = {}
    for rw in rows[hdr_i + 1:]:
        if not rw:
            continue
        tk = g(rw, "ticker")
        typ = str(g(rw, "security type") or "")
        if not tk or ("equity" not in typ.lower() and "stock" not in typ.lower() and "adr" not in typ.lower()):
            continue
        if not asof:
            eff = g(rw, "effective date")
            asof = eff.strftime("%Y-%m-%d") if hasattr(eff, "strftime") else iso(str(eff or ""))
        sh = num(g(rw, "shares/par"))
        if not sh:
            continue
        y = yahoo_sym(str(tk))
        out[y] = {"name": str(g(rw, "security description") or tk).title(), "shares": sh,
                  "weight": num(g(rw, "% of net assets")) or 0.0, "type": "equity", "yahoo": y}
    units = None
    try:
        pd_ = S.get(f"{JPM_BASE}/product-data", params={"cusip": cusip, "country": "us", "role": "adv",
                                                        "language": "en", "userLoggedIn": "false"}, timeout=30).json()
        units = num(((pd_.get("fundData") or pd_).get("shareClass") or {}).get("sharesOutstanding")) \
            if isinstance(pd_, dict) else None
        if units is None:
            units = _find_key(pd_, "sharesOutstanding")
    except Exception:
        pass
    return {"asof": asof, "units": units, "holdings": out}


def _find_key(obj, key):
    if isinstance(obj, dict):
        for k, v in obj.items():
            if k == key:
                return num(v)
            f = _find_key(v, key)
            if f is not None:
                return f
    elif isinstance(obj, list):
        for v in obj:
            f = _find_key(v, key)
            if f is not None:
                return f
    return None


# ---------------------------------------------------------------- Capital Group
def capital_group(ticker: str):
    url = f"https://www.capitalgroup.com/api/investments/investment-service/v1/etfs/{ticker}/holdings"
    r = S.get(url, params={"audience": "advisor"}, timeout=30)
    r.raise_for_status()
    d = r.json()
    dh = d.get("dailyHoldings") or {}
    asof = iso(dh.get("asOfDate", ""))
    out = {}
    for h in dh.get("holdings") or []:
        if not is_equity(h.get("assetClass")):
            continue
        tk = (h.get("ticker") or "").strip()
        sh = num(h.get("sharesOrPrincipalAmount"))
        if not tk or not sh:
            continue
        y = cg_yahoo(tk, h.get("isin"))
        out[y] = {"name": (h.get("securityName") or tk).title(), "shares": sh,
                  "weight": num(h.get("percentageOfNetAssets")) or 0.0, "type": "equity", "yahoo": y}
    return {"asof": asof, "units": None, "holdings": out}


ISIN_SUFFIX = {"JP": ".T", "GB": ".L", "HK": ".HK", "NL": ".AS", "FR": ".PA", "DE": ".DE", "CH": ".SW",
               "KR": ".KS", "TW": ".TW", "CA": ".TO", "IT": ".MI", "ES": ".MC", "SE": ".ST", "DK": ".CO",
               "AU": ".AX", "IN": ".NS", "NO": ".OL", "FI": ".HE", "BE": ".BR", "IE": "", "BR": ".SA"}


def cg_yahoo(tk: str, isin: str | None):
    tk = tk.strip().upper()
    if " " in tk:  # 'ASML NA' 같은 블룸버그 표기
        base, ex = tk.rsplit(" ", 1)
        m = {"NA": ".AS", "LN": ".L", "JP": ".T", "JT": ".T", "HK": ".HK", "FP": ".PA", "GR": ".DE", "GY": ".DE",
             "SW": ".SW", "SE": ".SW", "KS": ".KS", "TT": ".TW", "CN": ".TO", "IM": ".MI", "SM": ".MC",
             "SS": ".ST", "DC": ".CO", "AU": ".AX", "IN": ".NS", "US": "", "UN": "", "UW": "", "UQ": ""}
        suf = m.get(ex, "")
        if suf == ".HK":
            base = base.lstrip("0").zfill(4)
        return base.replace("/", "-").replace(".", "-") + suf
    cc = (isin or "")[:2]
    if cc and cc not in ("US",) and cc in ISIN_SUFFIX and not tk.isalpha() or (cc in ("JP", "HK", "KR", "TW")):
        suf = ISIN_SUFFIX.get(cc, "")
        if suf == ".HK":
            tk = tk.lstrip("0").zfill(4)
        return tk + suf
    return tk.replace("/", "-").replace(".", "-")


# ---------------------------------------------------------------- ARK
ARK_FILES = {"ARKK": "ARK_INNOVATION_ETF_ARKK_HOLDINGS.csv",
             "ARKW": "ARK_NEXT_GENERATION_INTERNET_ETF_ARKW_HOLDINGS.csv",
             "ARKG": "ARK_GENOMIC_REVOLUTION_ETF_ARKG_HOLDINGS.csv"}


def ark(ticker: str):
    url = "https://assets.ark-funds.com/fund-documents/funds-etf-csv/" + ARK_FILES[ticker]
    r = S.get(url, timeout=30)
    r.raise_for_status()
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig", errors="replace"))))
    out, asof = {}, None
    for row in rows:
        row = {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
        tk = row.get("ticker")
        sh = num(row.get("shares"))
        if not tk or not sh:
            continue
        asof = asof or iso(row.get("date", ""))
        y = yahoo_sym(tk)
        out[y] = {"name": row.get("company", tk).title(), "shares": sh,
                  "weight": num(row.get("weight (%)")) or 0.0, "type": "equity", "yahoo": y}
    # ARK 파일 날짜는 '다음 거래일' 표기 → 그대로 기준일로 사용(전일 종가 반영 보유)
    return {"asof": asof, "units": None, "holdings": out}


# ---------------------------------------------------------------- 일반 CSV/XLSX (Sequoia, GMO, Roundhill 등)
def _table_from_bytes(content: bytes):
    if content[:2] == b"PK":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    text = content.decode("utf-8-sig", errors="replace")
    return [r for r in csv.reader(io.StringIO(text))]


def generic_table(url: str, fund_filter: str | None = None):
    r = S.get(url, timeout=40)
    r.raise_for_status()
    rows = _table_from_bytes(r.content)
    hdr_i = None
    for i, rw in enumerate(rows[:40]):
        cells = [str(c).strip().lower() for c in rw if c is not None]
        if any(c in ("ticker", "symbol", "stockticker") for c in cells) and any("share" in c or "quantity" in c for c in cells):
            hdr_i = i
            break
    if hdr_i is None:
        raise ValueError("보유 표 헤더를 찾지 못함")
    hdr = [str(c).strip().lower() if c is not None else "" for c in rows[hdr_i]]

    def find(*keys):
        for k in keys:
            for i, h in enumerate(hdr):
                if h == k:
                    return i
        for k in keys:
            for i, h in enumerate(hdr):
                if k in h:
                    return i
        return None
    it, iname, ish, iw = find("ticker", "symbol", "stockticker"), find("name", "security", "description", "company"), \
        find("shares", "quantity", "shares/par"), find("weight", "% of net assets", "% net assets", "percent")
    idate, ifund = find("date", "as of"), find("account", "fund")
    asof = None
    for rw in rows[:hdr_i]:
        for c in rw:
            m = re.search(r"\d{1,2}/\d{1,2}/\d{2,4}", str(c or ""))
            if m:
                asof = iso(m.group(0))
    out = {}
    for rw in rows[hdr_i + 1:]:
        if not rw or it is None or it >= len(rw):
            continue
        if fund_filter and ifund is not None and str(rw[ifund] or "").strip().upper() != fund_filter:
            continue
        tk = str(rw[it] or "").strip()
        sh = num(rw[ish]) if ish is not None and ish < len(rw) else None
        if not tk or not sh or tk.upper() in ("CASH", "USD", "FGXXX", "MONEY MARKET"):
            continue
        if idate is not None and not asof:
            v = rw[idate]
            asof = v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else iso(str(v or ""))
        y = cg_yahoo(tk, None)
        out[y] = {"name": str(rw[iname] or tk).title() if iname is not None else tk, "shares": sh,
                  "weight": (num(rw[iw]) or 0.0) if iw is not None and iw < len(rw) else 0.0,
                  "type": "equity", "yahoo": y}
    if out and all(h["weight"] <= 1.0 for h in out.values()) and sum(h["weight"] for h in out.values()) <= 1.5:
        for h in out.values():
            h["weight"] *= 100  # 0.05 → 5%
    return {"asof": asof, "units": None, "holdings": out}


def sequoia():
    page = S.get("https://sequoia.filepoint.live/sequoia-etf", timeout=30).text
    m = re.findall(r'(assets/data/Sequoia_ETF_Holdings_[\d-]+\.csv)', page)
    if not m:
        raise ValueError("Sequoia 보유 CSV 링크를 찾지 못함")
    return generic_table("https://sequoia.filepoint.live/" + sorted(set(m))[-1])


def gmo_qlty():
    return generic_table("https://www.gmo.com/globalassets/documents---manually-loaded/documents/u.s.-quality-etf_etf_holdings")


def roundhill(ticker: str, d: datetime):
    url = f"https://www.roundhillinvestments.com/assets/data/FilepointRoundhill.40RU.RU_Holdings_{d:%m%d%Y}.csv"
    return generic_table(url, fund_filter=ticker)
